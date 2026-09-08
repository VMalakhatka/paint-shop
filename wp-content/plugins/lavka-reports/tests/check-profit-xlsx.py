"""Independently verify browser-generated workbooks with openpyxl and ZIP CRCs."""
import json
import sys
import zipfile
from pathlib import Path
import openpyxl

folder = Path(sys.argv[1])
fixture = json.loads((Path(__file__).parent / 'fixtures/profit-report.json').read_text())
for name in ('profit.xlsx', 'java-contract.xlsx'):
    path = folder / name
    with zipfile.ZipFile(path) as archive:
        assert archive.testzip() is None, name
        assert not any(b'<f>' in archive.read(member) for member in archive.namelist() if member.endswith('.xml'))
    book = openpyxl.load_workbook(path)
    assert len(book.sheetnames) == 12
    assert all(sheet.freeze_panes == 'A4' for sheet in book)
    if name == 'profit.xlsx':
        kyiv = book['Kyiv']
        headers = [cell.value for cell in kyiv[3]]
        zero = kyiv.cell(5, headers.index('Amount') + 1)
        assert zero.value == 0 and zero.data_type == 'n'
        assert book['Odesa'].max_row == 4  # City filter must not remove Odesa from XLSX.
        audit = book['Document audit']
        headers = [cell.value for cell in audit[3]]
        identity = audit.cell(4, headers.index('Payment ID') + 1)
        assert identity.value == '9007199254740993' and identity.data_type == 's'
        literal = audit.cell(4, headers.index('Document number') + 1)
        assert literal.value.startswith('=HYPERLINK') and literal.data_type == 's'
        assert audit.cell(4, headers.index('Report amount, UAH') + 1).data_type == 'n'
        assert 'truncat' in audit.cell(2, 1).value.lower() or '500' in audit.cell(2, 1).value
        assert book['Period checks'].max_row == 4
    else:
        assert book['All expense rows'].max_row == len(fixture['expenseLines']) + 3
        assert book['Document audit'].max_row == len(fixture['documents']) + 3
        assert book['Period checks'].max_row == len(fixture['periodDiagnostics']) + 3
        for city, sheet_name in [('KYIV', 'Kyiv'), ('ODESA', 'Odesa')]:
            sheet = book[sheet_name]
            headers = [cell.value for cell in sheet[3]]
            amount_col = headers.index('Amount') + 1
            id_col = headers.index('Expense row ID') + 1
            observed = {sheet.cell(n, id_col).value: sheet.cell(n, amount_col).value for n in range(4, sheet.max_row + 1)}
            expected = {row['lineId']: float(row['amount']) for row in fixture['expenseLines'] if row['city'] == city}
            assert observed == expected, city
    print('PASS:', name, 'CRC, schema, coverage, precision and literal text')

partial = openpyxl.load_workbook(folder / 'partial.xlsx')
assert 'Report section availability' in partial.sheetnames
section_rows = list(partial['Report section availability'].values)
assert any('UNAVAILABLE' in row for row in section_rows)
profit = partial['Profit by city']
headers = [cell.value for cell in profit[3]]
assert profit.cell(5, headers.index('Profit') + 1).value == '—'
assert partial['Kyiv'].max_row > 3
print('PASS: partial.xlsx availability and unknown dependent profit preserved')

partial_java = openpyxl.load_workbook(folder / 'partial-java.xlsx')
assert 'Report section availability' in partial_java.sheetnames
profit = partial_java['Profit by city']
headers = [cell.value for cell in profit[3]]
assert all(profit.cell(row, headers.index('Profit') + 1).value == '—' for row in [4, 5])
assert partial_java['Control totals'].max_row <= 3
assert not any(cell.data_type == 'n' and cell.value is not None for row in partial_java['Control totals'] for cell in row)
assert partial_java['Inventory accounting value'].max_row > 3
print('PASS: partial-java.xlsx failed expenses preserve unknown totals and available inventory')
